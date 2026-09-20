"""
server.py  (slim build)
========================
Local Flask server: upload an AR PDF, run ONLY Step 1 (trimmed core-financial
PDF) + Step 2 (Excel workbook) of the pipeline, then let the user view the
Excel output in-browser and download both output files.

This mirrors the look-and-feel of the original repo's server.py (upload
form -> live progress terminal -> result page with an Excel viewer) but the
background job runs `extract_core_financial_statements` + `extract_financials`
directly instead of the full memo/red-flag/narrative agent graph, and there
is no memo to render.

Usage:
    python server.py
    -> open http://127.0.0.1:5000
"""
import os, sys, traceback, threading, uuid, io, logging
from pathlib import Path
from flask import (
    Flask, request, render_template_string, redirect, url_for, session,
    jsonify, send_file,
)
import openpyxl

_ROOT = str(Path(__file__).resolve().parent)
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

from pipeline.Step1 import extract_core_financial_statements
from pipeline.Step2 import extract_financials
from pipeline.taxonomy import TAXONOMY, Sign, Statement

UPLOAD_DIR = Path(_ROOT) / "uploads"
OUTPUT_DIR = Path(_ROOT) / "output"
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "ar-due-diligence-dev-key")

_IMPORTANT_NODES = {
    "REVENUE_FROM_OPERATIONS", "TOTAL_INCOME", "EBITDA", "PROFIT_BEFORE_TAX",
    "PROFIT_FOR_THE_YEAR", "TOTAL_COMPREHENSIVE_INCOME", "EARNINGS_PER_SHARE",
    "TOTAL_ASSETS", "TOTAL_EQUITY", "TOTAL_LIABILITIES", "TOTAL_EQUITY_AND_LIABILITIES",
    "TOTAL_NON_CURRENT_ASSETS", "TOTAL_CURRENT_ASSETS", "TOTAL_NON_CURRENT_LIABILITIES",
    "TOTAL_CURRENT_LIABILITIES", "NET_CASH_FROM_OPERATING", "NET_CASH_FROM_INVESTING",
    "NET_CASH_FROM_FINANCING", "NET_CHANGE_IN_CASH", "CLOSING_CASH_BALANCE",
    "CASH_AND_CASH_EQUIVALENTS", "GROSS_FIXED_ASSETS", "NET_FIXED_ASSETS",
    "NET_WORTH", "GROSS_DEBT", "RESERVE_CLOSING_BALANCE",
}

# ── Row category → background colour (muted, single blue-grey family) ──────
_CATEGORY_COLORS = {
    "heading":    "#dce1f0",
    "total":      "#c9d6ee",
    "income":     "#e4ecf7",
    "expense":    "#eef1f8",
    "asset":      "#e7edf6",
    "liability":  "#eef0f7",
    "cashflow":   "#e9edf6",
    "unmapped":   "#f5f6fa",
}

def _row_category(node_val: str, is_heading: bool) -> str:
    if is_heading:
        return "heading"
    node = TAXONOMY.get(node_val) if node_val else None
    if node is None:
        return "unmapped"
    if node.is_total:
        return "total"
    if node.statement == Statement.CASH_FLOW:
        return "cashflow"
    if node.statement == Statement.BALANCE_SHEET:
        return "liability" if node.sign == Sign.NEGATIVE else "asset"
    return "expense" if node.sign == Sign.NEGATIVE else "income"


# ── Background job tracking ─────────────────────────────────────────────────
_JOBS = {}
_JOBS_LOCK = threading.Lock()

_STAGES = ["EXTRACT CORE PAGES", "PARSE & BUILD EXCEL"]


class _JobStream(io.TextIOBase):
    """Captures stdout / logging output line-by-line into a job's live log."""
    def __init__(self, job_id):
        self.job_id = job_id
        self._buf = ""

    def write(self, s):
        self._buf += s
        while "\n" in self._buf:
            line, self._buf = self._buf.split("\n", 1)
            if line.strip():
                _push_line(self.job_id, line.strip())
        return len(s)

    def flush(self):
        pass


def _push_line(job_id, line):
    with _JOBS_LOCK:
        job = _JOBS.get(job_id)
        if not job:
            return
        job["lines"].append(line)


def _run_job(job_id, pdf_path, ticker):
    stream = _JobStream(job_id)
    old_stdout = sys.stdout
    sys.stdout = stream

    # Step1 logs via the `logging` module (log.info/.warning/.error) rather
    # than print(); route the root logger into the same terminal stream so
    # Phase 1/2/3 progress shows up live too.
    log_handler = logging.StreamHandler(stream)
    log_handler.setFormatter(logging.Formatter("%(message)s"))
    root_logger = logging.getLogger()
    old_level = root_logger.level
    root_logger.addHandler(log_handler)
    root_logger.setLevel(logging.INFO)

    try:
        gemini_key = os.environ.get("GEMINI_API_KEY", "")
        trimmed_pdf = str(OUTPUT_DIR / f"{ticker}_trimmed.pdf")
        output_xlsx = str(OUTPUT_DIR / f"{ticker}.xlsx")

        print(f"── Step 1: Extracting core financial pages → {trimmed_pdf}")
        with _JOBS_LOCK:
            _JOBS[job_id]["stage_idx"] = 0
        extract_core_financial_statements(pdf_path, trimmed_pdf, gemini_key)

        print(f"── Step 2: Parsing & building Excel → {output_xlsx}")
        with _JOBS_LOCK:
            _JOBS[job_id]["stage_idx"] = 1
        extract_financials(trimmed_pdf, output_xlsx)

        with _JOBS_LOCK:
            _JOBS[job_id]["trimmed_pdf"] = trimmed_pdf
            _JOBS[job_id]["xlsx_path"] = output_xlsx
            _JOBS[job_id]["stage_idx"] = len(_STAGES) - 1
            _JOBS[job_id]["done"] = True
    except Exception as exc:
        traceback.print_exc()
        with _JOBS_LOCK:
            _JOBS[job_id]["error"] = str(exc)
            _JOBS[job_id]["done"] = True
    finally:
        root_logger.removeHandler(log_handler)
        root_logger.setLevel(old_level)
        sys.stdout = old_stdout


PAGE = """
<!doctype html>
<html><head>
<meta charset="utf-8">
<title>AR Due Diligence — Extractor</title>
<style>
  body { font-family: 'Segoe UI', sans-serif; max-width: 860px; margin: 40px auto; padding: 0 20px; background: #f7f8fb; color: #1f2430; }
  h1 { color: #2b2f77; }
  form { background: #fff; padding: 24px; border-radius: 12px; box-shadow: 0 2px 10px rgba(0,0,0,.08); }
  input[type=text], input[type=file] { padding: 8px; margin: 8px 0 16px; width: 100%; box-sizing: border-box; border: 1px solid #ccc; border-radius: 6px; }
  button { background: #4f5bd5; color: #fff; border: none; padding: 10px 22px; border-radius: 8px; font-size: 15px; cursor: pointer; }
  button:hover { background: #3c46b0; }
  .result { background: #fff; padding: 32px; border-radius: 12px; box-shadow: 0 2px 10px rgba(0,0,0,.08); margin-top: 24px; }
  .error { background: #ffe9e9; color: #a30000; padding: 16px; border-radius: 8px; }
  .actions { margin-top: 16px; display:flex; gap:10px; flex-wrap: wrap; }
  .btn-secondary { background:#00b39f; color:#fff; border:none; padding:10px 22px; border-radius:8px; font-size:15px; cursor:pointer; text-decoration:none; display:inline-block; }
  .btn-secondary:hover { background:#00937f; }
  .btn-tertiary { background:#4f5bd5; }
  .btn-tertiary:hover { background:#3c46b0; }
</style>
</head><body>
<h1>📄 AR Financial Extractor</h1>
<p>Upload an annual report PDF — this extracts the core financial-statement pages and builds the Excel workbook (Step 1 + Step 2 only).</p>
<form method="post" action="{{ url_for('start') }}" enctype="multipart/form-data">
  <label>Annual Report PDF</label>
  <input type="file" name="pdf" accept="application/pdf" required>
  <label>Ticker / company code, e.g. TCS</label>
  <input type="text" name="ticker" required>
  <button type="submit">Run Extraction</button>
</form>
{% if error %}
<div class="error"><b>Error:</b> {{ error }}</div>
{% endif %}
{% if xlsx_path %}
<div class="result">
  <h2>✅ Done</h2>
  <p>Core financial pages extracted and Excel workbook built.</p>
  <div class="actions">
    <a class="btn-secondary btn-tertiary" href="{{ url_for('view_excel', path=xlsx_path) }}">📈 View Excel Output</a>
    <a class="btn-secondary" href="{{ url_for('download_file', path=xlsx_path) }}">⬇ Download Excel</a>
    {% if trimmed_pdf %}
    <a class="btn-secondary" href="{{ url_for('download_file', path=trimmed_pdf) }}">⬇ Download Trimmed PDF</a>
    {% endif %}
    <a class="btn-secondary" style="background:#888" href="{{ url_for('index', reset=1) }}">🔄 New Extraction</a>
  </div>
</div>
{% endif %}
</body></html>
"""

EXCEL_PAGE = """
<!doctype html>
<html><head>
<meta charset="utf-8">
<title>Excel Output</title>
<style>
  body { font-family: 'Segoe UI', sans-serif; max-width: 1100px; margin: 40px auto; padding: 0 20px; background: #f7f8fb; color: #1f2430; }
  h1 { color: #2b2f77; }
  a.back { color:#4f5bd5; text-decoration:none; font-weight:600; }
  .card { background:#fff; padding:24px; border-radius:12px; box-shadow:0 2px 10px rgba(0,0,0,.08); margin-top:16px; }
  .tabs { display:flex; gap:6px; flex-wrap:wrap; margin-bottom:14px; }
  .tab-btn { background:#eef0fb; border:none; padding:8px 16px; border-radius:20px; cursor:pointer; font-weight:600; color:#2b2f77; }
  .tab-btn.active { background:#4f5bd5; color:#fff; }
  .sheet-table { display:none; overflow-x:auto; }
  .sheet-table.active { display:block; }
  table.xl { border-collapse:collapse; width:100%; font-size:13px; }
  table.xl th { background:#4f5bd5; color:#fff; padding:8px 10px; position:sticky; top:0; }
  table.xl td { padding:6px 10px; border:1px solid #e2e4f0; }
  table.xl tr:nth-child(even) td { background:#f4f5fc; }
  table.xl tr:hover td { background:#e9ecfc; }
  .num { text-align:right; color:#1a3a8f; font-variant-numeric: tabular-nums; }
  .important-row td { background:#fff6da !important; font-weight:700; }
  .important-row td.num.important { color:#b8860b; }
  td.important:not(.num) { color:#1a3a1a; }
</style>
<script>
function showSheet(name){
  document.querySelectorAll('.sheet-table').forEach(el=>el.classList.remove('active'));
  document.querySelectorAll('.tab-btn').forEach(el=>el.classList.remove('active'));
  document.getElementById('sheet-'+name).classList.add('active');
  document.getElementById('tab-'+name).classList.add('active');
}
</script>
</head><body>
<a class="back" href="{{ url_for('index') }}">&larr; Back</a>
<h1>📈 Excel Output</h1>
<div class="card">
  <div class="tabs">
    {% for name in sheet_names %}
    <button class="tab-btn {{ 'active' if loop.first else '' }}" id="tab-{{ name|replace(' ','_') }}" onclick="showSheet('{{ name|replace(' ','_') }}')">{{ name }}</button>
    {% endfor %}
  </div>
  {% for name, table_html in sheets %}
  <div class="sheet-table {{ 'active' if loop.first else '' }}" id="sheet-{{ name|replace(' ','_') }}">{{ table_html | safe }}</div>
  {% endfor %}
</div>
</body></html>
"""

PROGRESS_PAGE = """
<!doctype html>
<html><head>
<meta charset="utf-8">
<title>Extracting…</title>
<style>
  body { font-family: 'Segoe UI', sans-serif; max-width: 780px; margin: 40px auto; padding: 0 20px; background: #0e1020; color: #d9e0ff; }
  h1 { color: #8fa2ff; text-align:center; }
  .stages { display:flex; justify-content:space-between; margin: 28px 0 18px; }
  .stage { flex:1; text-align:center; position:relative; }
  .stage .dot { width:34px; height:34px; border-radius:50%; background:#262a4d; border:2px solid #444a80; margin:0 auto 8px; display:flex; align-items:center; justify-content:center; font-weight:700; transition: all .4s ease; }
  .stage.active .dot { background:#4f5bd5; border-color:#8fa2ff; box-shadow:0 0 18px #4f5bd5; animation: pulse 1.2s infinite; }
  .stage.done .dot { background:#00b39f; border-color:#00e0c0; box-shadow:0 0 12px #00b39f; }
  .stage span.label { font-size:12px; color:#9aa4d6; }
  .stage.active span.label { color:#c9d2ff; font-weight:700; }
  .bar-track { height:6px; background:#20223e; border-radius:6px; overflow:hidden; margin-bottom:26px; }
  .bar-fill { height:100%; width:0%; background:linear-gradient(90deg,#4f5bd5,#00b39f); transition: width .6s ease; }
  @keyframes pulse { 0%{ transform:scale(1);} 50%{ transform:scale(1.15);} 100%{ transform:scale(1);} }
  .terminal { background:#05060f; border:1px solid #262a4d; border-radius:10px; padding:18px 20px; height:340px; overflow-y:auto; font-family:'Consolas','Menlo',monospace; font-size:13px; line-height:1.6; box-shadow: inset 0 0 30px rgba(79,91,213,.08); }
  .terminal .ln { opacity:0; animation: fadeIn .35s forwards; white-space:pre-wrap; }
  .terminal .ln.ok { color:#5be08a; }
  .terminal .ln.warn { color:#ffcf5c; }
  .terminal .ln.plain { color:#9aa4d6; }
  @keyframes fadeIn { from{opacity:0; transform:translateY(4px);} to{opacity:1; transform:translateY(0);} }
  .errbox { background:#3a0d16; color:#ff8f8f; padding:14px 18px; border-radius:8px; margin-top:16px; }
</style>
</head><body>
<h1>⚙️ Extracting…</h1>
<div class="stages" id="stages">
  {% for s in stages %}
  <div class="stage" id="stage-{{ loop.index0 }}">
    <div class="dot">{{ loop.index }}</div>
    <span class="label">{{ s }}</span>
  </div>
  {% endfor %}
</div>
<div class="bar-track"><div class="bar-fill" id="barFill"></div></div>
<div class="terminal" id="terminal"></div>
<div id="errWrap"></div>

<script>
const jobId = "{{ job_id }}";
const totalStages = {{ stages|length }};
let shownLines = 0;

function classify(line){
  if (line.startsWith('──')) return 'node';
  if (line.includes('✅') || line.toLowerCase().includes('complete')) return 'ok';
  if (line.includes('⚠️') || line.toLowerCase().includes('warn')) return 'warn';
  return 'plain';
}

function updateStages(stageIdx, done){
  for (let i = 0; i < totalStages; i++){
    const el = document.getElementById('stage-' + i);
    el.classList.remove('active','done');
    if (i < stageIdx || (done && i <= stageIdx)) el.classList.add('done');
    else if (i === stageIdx) el.classList.add('active');
  }
  const pct = done ? 100 : Math.min(95, ((stageIdx + 0.5) / totalStages) * 100);
  document.getElementById('barFill').style.width = pct + '%';
}

async function poll(){
  try {
    const res = await fetch('/status/' + jobId);
    const data = await res.json();
    const term = document.getElementById('terminal');
    for (; shownLines < data.lines.length; shownLines++){
      const div = document.createElement('div');
      div.className = 'ln ' + classify(data.lines[shownLines]);
      div.textContent = data.lines[shownLines];
      term.appendChild(div);
    }
    term.scrollTop = term.scrollHeight;
    updateStages(data.stage_idx, data.done);

    if (data.error){
      document.getElementById('errWrap').innerHTML = '<div class="errbox"><b>Error:</b> ' + data.error + '</div>';
      return;
    }
    if (data.done){
      setTimeout(() => { window.location.href = '/result/' + jobId; }, 700);
      return;
    }
    setTimeout(poll, 600);
  } catch(e){
    setTimeout(poll, 1200);
  }
}
poll();
</script>
</body></html>
"""

@app.route("/", methods=["GET"])
def index():
    if request.args.get("reset"):
        session.pop("xlsx_path", None)
        session.pop("trimmed_pdf", None)
        return redirect(url_for("index"))
    return render_template_string(
        PAGE,
        error=None,
        xlsx_path=session.get("xlsx_path"),
        trimmed_pdf=session.get("trimmed_pdf"),
    )


@app.route("/start", methods=["POST"])
def start():
    f = request.files.get("pdf")
    ticker = request.form.get("ticker", "").strip().upper()
    if not f or not f.filename.lower().endswith(".pdf") or not ticker:
        return render_template_string(PAGE, error="Please provide a valid PDF and ticker.", xlsx_path=None, trimmed_pdf=None)

    if not os.environ.get("GEMINI_API_KEY"):
        return render_template_string(PAGE, error="GEMINI_API_KEY is not set on the server.", xlsx_path=None, trimmed_pdf=None)

    pdf_path = UPLOAD_DIR / f.filename
    f.save(pdf_path)

    job_id = uuid.uuid4().hex
    with _JOBS_LOCK:
        _JOBS[job_id] = {"lines": [], "stage_idx": 0, "done": False, "error": None,
                          "xlsx_path": None, "trimmed_pdf": None}

    t = threading.Thread(target=_run_job, args=(job_id, str(pdf_path), ticker), daemon=True)
    t.start()

    return redirect(url_for("progress_page", job_id=job_id))


@app.route("/progress/<job_id>")
def progress_page(job_id):
    if job_id not in _JOBS:
        return redirect(url_for("index"))
    return render_template_string(PROGRESS_PAGE, job_id=job_id, stages=_STAGES)


@app.route("/status/<job_id>")
def status(job_id):
    with _JOBS_LOCK:
        job = _JOBS.get(job_id)
        if not job:
            return jsonify({"lines": [], "stage_idx": 0, "done": True, "error": "Job not found"})
        return jsonify({
            "lines": job["lines"],
            "stage_idx": job["stage_idx"],
            "done": job["done"],
            "error": job["error"],
        })


@app.route("/result/<job_id>")
def result(job_id):
    with _JOBS_LOCK:
        job = _JOBS.get(job_id)
    if not job:
        return redirect(url_for("index"))
    session["xlsx_path"] = job.get("xlsx_path")
    session["trimmed_pdf"] = job.get("trimmed_pdf")
    with _JOBS_LOCK:
        _JOBS.pop(job_id, None)
    return redirect(url_for("index"))


def _cell_html(value, bold=False, bg=None):
    cls = "num" if isinstance(value, (int, float)) else ""
    cls = (cls + " important").strip() if bold else cls
    cls_attr = f' class="{cls}"' if cls else ""
    style_attr = f' style="background:{bg};"' if bg else ""
    if isinstance(value, float):
        return f'<td{cls_attr}{style_attr}>{value:,.2f}</td>'
    if isinstance(value, int):
        return f'<td{cls_attr}{style_attr}>{value:,}</td>'
    return f"<td{cls_attr}{style_attr}>{'' if value is None else value}</td>"


@app.route("/excel")
def view_excel():
    path = request.args.get("path", "")
    if not path or not Path(path).exists():
        return render_template_string(PAGE, error="Excel file not found.", xlsx_path=None, trimmed_pdf=None)

    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_names = wb.sheetnames
    sheets = []
    for name in sheet_names:
        ws = wb[name]
        rows_html = []
        header = None
        taxonomy_col = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            row = list(row)
            if i == 0:
                header = row
                taxonomy_col = next(
                    (idx for idx, h in enumerate(header) if h and "taxonomy" in str(h).lower()),
                    None,
                )
                display_row = [v for idx, v in enumerate(row) if idx != taxonomy_col]
                cells = "".join(f"<th>{'' if v is None else v}</th>" for v in display_row)
                rows_html.append(f"<tr>{cells}</tr>")
                continue

            node_val = str(row[taxonomy_col]).strip().upper() if taxonomy_col is not None and taxonomy_col < len(row) else ""
            is_important = node_val in _IMPORTANT_NODES
            display_row = [v for idx, v in enumerate(row) if idx != taxonomy_col]
            numeric_vals = [v for idx, v in enumerate(row)
                             if idx != taxonomy_col and idx != 0
                             and isinstance(v, (int, float))]
            is_heading = len(numeric_vals) == 0
            category = _row_category(node_val, is_heading)
            row_color = _CATEGORY_COLORS[category]
            cells = "".join(_cell_html(v, bold=is_important, bg=None if is_important else row_color) for v in display_row)
            row_cls = ' class="important-row"' if is_important else ""
            rows_html.append(f"<tr{row_cls}>{cells}</tr>")

        table_html = f'<table class="xl">{"".join(rows_html)}</table>'
        sheets.append((name, table_html))

    return render_template_string(EXCEL_PAGE, sheet_names=sheet_names, sheets=sheets)


@app.route("/download")
def download_file():
    path = request.args.get("path", "")
    if not path or not Path(path).exists():
        return render_template_string(PAGE, error="File not found.", xlsx_path=None, trimmed_pdf=None)
    return send_file(path, as_attachment=True)


if __name__ == "__main__":
    missing = [k for k in ["GEMINI_API_KEY"] if not os.environ.get(k)]
    if missing:
        print(f"⚠️  Missing env vars: {missing} — add to .env before running.")
    app.run(host="127.0.0.1", port=5000, debug=True)